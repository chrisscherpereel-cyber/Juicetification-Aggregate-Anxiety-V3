"""Safety stock, beginning backorders and the worker policy: stated clearly, applied the
same way by the answer key, the workbook formulas (calculated by real Excel) and the checker."""

import io
import math

import pytest
from openpyxl import load_workbook

from aggplan import checker, engine, hints, layout as L, text, workbook
from tests.helpers import SID, fixture, scn7
from tests.make_fixtures import POLICY_CFG


@pytest.mark.parametrize("piece", ["chase", "level", "hybrid", "summary"])
def test_excel_answer_key_with_every_policy_switched(piece):
    res = checker.check_piece(fixture("policies_seed7.xlsx"), scn7(**POLICY_CFG), piece, SID)
    assert res.complete, (res.notes, res.messages, res.counts(), res.feasibility)


def test_beginning_backorders_are_given_data_with_a_name():
    s = scn7(beginning_backorders=500)
    wb = load_workbook(io.BytesIO(workbook.build_workbook(s, student_id=SID)))
    assert "BeginBackorders" in wb.defined_names
    assert any(r[0].startswith("Beginning backorders") and r[1] == "500" for r in text.param_rows(s))
    df = engine.reference_plan(s, "chase")
    assert df["Beginning Backorders"].iloc[0] == 500
    assert df["Regular Production"].iloc[0] == s.demand[0] - s["beginning_inventory"] + 500


def test_level_rate_pays_back_orders_owed():
    base, owed = scn7(), scn7(beginning_backorders=1200)
    assert engine.level_rate(owed) == engine.level_rate(base) + 100


def test_partial_workers_keep_fractions_whole_workers_round_up():
    whole, part = scn7(), scn7(worker_policy="partial")
    assert all(float(w).is_integer() for w in engine.reference_plan(whole, "chase")["Workers"])
    w = engine.reference_plan(part, "chase")["Workers"]
    assert any(not float(x).is_integer() for x in w)
    assert math.isclose(w.iloc[0], engine.reference_plan(part, "chase")["Regular Production"].iloc[0]
                        / part.cap_month)
    labels = [c["label"] for c in engine.check_feasibility(engine.reference_plan(part, "chase"), part)]
    assert "Whole-number workforce" not in labels


def test_policies_are_stated_for_every_option():
    for cfg in ({}, POLICY_CFG):
        s = scn7(**cfg)
        titles = [t for t, _, _ in text.policies(s)]
        assert [t.split(":")[0] for t in titles] == ["Safety stock", "Backorders", "Workforce"]
        body = " ".join(" ".join([sm] + d) for _, sm, d in text.policies(s))
        assert "EndTarget" in body and "BeginBackorders" in body and "StartWorkers" in body
    assert "NONE" in text.policies(scn7())[0][0] and "whole" in text.policies(scn7())[2][0]
    assert "KEEP" in text.policies(scn7(**POLICY_CFG))[0][0]
    assert "fractional" in text.policies(scn7(**POLICY_CFG))[2][0]


def test_policies_appear_on_scenario_and_start_sheets():
    wb = load_workbook(io.BytesIO(workbook.build_workbook(scn7(**POLICY_CFG), student_id=SID)))
    for sheet in (L.SHEET_SCENARIO, L.SHEET_START):
        cells = " ".join(str(c.value) for row in wb[sheet].iter_rows() for c in row if c.value)
        assert "Safety stock: KEEP" in cells and "fractional workers allowed" in cells
        assert "BeginBackorders" in cells


def test_hints_follow_the_worker_policy():
    assert "ROUNDUP" in hints.ladder("chase", "Workers", True)[2]
    assert "ROUNDUP" not in " ".join(hints.ladder("chase", "Workers", False))
    assert "ROUNDUP" not in " ".join(hints.ladder("level", "plan:level_workers", False))
    assert "BeginBackorders" in hints.ladder("chase", "Beginning Backorders")[2]
