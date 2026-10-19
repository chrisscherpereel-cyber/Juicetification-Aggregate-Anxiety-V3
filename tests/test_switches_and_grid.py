"""Safety stock hidden when unused, the policy switches, and the full cell-by-cell grid."""

import io
import os

import pytest
from openpyxl import load_workbook

from aggplan import checker, layout as L, text, workbook
from aggplan.checker import Cell
from tests.helpers import SID, book_from, fixture, scn7

APP = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "app.py")


def _cells(wb, sheet):
    return " ".join(str(c.value) for row in wb[sheet].iter_rows() for c in row if c.value)


def test_safety_stock_hidden_when_policy_is_none():
    s = scn7()
    wb = load_workbook(io.BytesIO(workbook.build_workbook(s, student_id=SID)))
    assert "SafetyStock" not in wb.defined_names and "EndTarget" in wb.defined_names
    assert "SafetyStock" not in _cells(wb, L.SHEET_SCENARIO)
    assert "2,400 bottles" not in " ".join(f"{a} {b}" for a, b, _ in text.param_rows(s)
                                           if "afety" in a)
    kept = scn7(safety_stock_policy="maintain")
    wb = load_workbook(io.BytesIO(workbook.build_workbook(kept, student_id=SID)))
    assert "SafetyStock" in wb.defined_names


def test_workbook_for_other_policies_is_explained():
    res = checker.check_piece(fixture("key_seed7.xlsx"), scn7(worker_policy="partial"), "chase", SID)
    assert res.fatal and "different policies" in res.fatal[0] and "workforce" in res.fatal[0]


def test_hybrid_grid_shows_every_column_before_decisions():
    book, _ = checker.read_book(fixture("key_seed7.xlsx"))
    b = book_from(book, {(L.SHEET_HYBRID, f"E{r}"): Cell(None, None) for r in range(2, 14)})
    res = checker.check_book(b, scn7(), "hybrid", SID)
    g = checker.grid(res)
    assert len(g.columns) == 20                      # every column the student built
    statuses = {c.status for c in res.cells if c.row_label != "Planning"}
    assert "waiting" in statuses and "blank" in statuses and not res.complete


def test_helper_columns_are_shown_but_not_graded():
    book, _ = checker.read_book(fixture("key_seed7.xlsx"))
    b = book_from(book, {(L.SHEET_CHASE, "R1"): Cell("Notes", None),
                         (L.SHEET_CHASE, "R2"): Cell(42.0, None)})
    res = checker.check_book(b, scn7(), "chase", SID)
    assert "R · Notes" in checker.grid(res).columns
    assert res.complete


st_testing = pytest.importorskip("streamlit.testing.v1")


def test_student_sets_the_two_policies_until_a_piece_is_submitted():
    at = st_testing.AppTest.from_file(APP, default_timeout=30).run()
    at.text_input[0].input("Ann")
    at.text_input[1].input("S1")
    next(b for b in at.button if b.label == "Start").click()
    at = at.run()
    assert [r.label for r in at.radio] == ["**Safety stock**", "**Workforce**"]
    at.radio[0].set_value("maintain")
    at.radio[1].set_value("partial")
    at = at.run()
    assert at.session_state["policy"] == {"safety_stock_policy": "maintain",
                                          "worker_policy": "partial"}
    assert not at.exception
    at.session_state["pieces"]["chase"]["submitted"] = {"at": "now", "sha": "x", "file": "f",
                                                        "metrics": {}, "feasibility": []}
    at = at.run()
    assert all(r.disabled for r in at.radio)
