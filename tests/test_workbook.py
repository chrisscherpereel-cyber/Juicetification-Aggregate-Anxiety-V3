import io

from openpyxl import load_workbook

from aggplan import layout as L, workbook
from tests.helpers import scn7


def _wb(cfg=None, **kw):
    s = scn7(**(cfg or {}))
    data = workbook.build_workbook(s, student_name="Ann", student_id="S1", **kw)
    return s, load_workbook(io.BytesIO(data))


def test_sheets_names_and_meta():
    s, wb = _wb()
    assert wb.sheetnames == [L.SHEET_START, L.SHEET_SCENARIO, L.SHEET_CHASE, L.SHEET_LEVEL,
                             L.SHEET_HYBRID, L.SHEET_SUMMARY, L.SHEET_HELP, L.SHEET_META]
    meta_ws = wb[L.SHEET_META]
    assert meta_ws.sheet_state == "veryHidden"
    meta = {meta_ws.cell(r, 1).value: meta_ws.cell(r, 2).value for r in range(1, 11)}
    assert meta["seed"] == s.seed and meta["config_hash"] == s.config_hash
    for name in ("LaborCost", "BeginInv", "EndTarget", "BottlesPerHour", "SubCapacity"):
        assert name in wb.defined_names


def test_students_design_the_columns():
    """Only Month and Demand are given; headings, formulas and totals are the student's."""
    s, wb = _wb()
    ws = wb[L.SHEET_CHASE]
    assert ws.protection.sheet
    assert ws["A1"].value == "Month" and ws["B1"].value == "Demand"
    for j in range(L.FIRST_STUDENT_COL, L.LAST_STUDENT_COL + 1):
        for r in [1] + [L.month_row(i) for i in range(12)] + [L.TOTAL_ROW]:
            c = ws.cell(row=r, column=j)
            assert c.value is None and not c.protection.locked
    for i in range(12):
        given = ws[f"B{L.month_row(i)}"]
        assert given.value == s.demand[i] and given.protection.locked
    for pc in L.PLANNING["level"]:
        c = wb[L.SHEET_LEVEL][L.planning_ref("level", pc.key)]
        assert c.value is None and not c.protection.locked


def test_no_rules_or_formulas_are_given_away():
    _, wb = _wb()
    for name in (L.SHEET_START, L.SHEET_CHASE, L.SHEET_LEVEL, L.SHEET_HYBRID, L.SHEET_SUMMARY,
                 L.SHEET_HELP):
        for row in wb[name].iter_rows():
            for c in row:
                v = str(c.value or "")
                assert "ROUNDUP(Production" not in v and "MAX(0, Demand" not in v, (name, c.coordinate)
                assert "Beginning Inventory +" not in v, (name, c.coordinate)


def test_hybrid_sheet_optional():
    _, wb = _wb(cfg={"include_hybrid": False})
    assert L.SHEET_HYBRID not in wb.sheetnames


def test_solution_writes_headings_and_formulas():
    _, wb = _wb(solution=True)
    ws = wb[L.SHEET_LEVEL]
    assert ws["C1"].value == "Beginning Inventory" and ws["E1"].value == "Production"
    assert ws["E5"].value == "=$C$20"
    assert wb[L.SHEET_SUMMARY][L.LOWER_COST_CELL].value.startswith("=IF(")
    assert wb[L.SHEET_SUMMARY][L.RECOMMEND_CELL].value == "Chase"


def test_student_text_is_never_a_formula():
    data = workbook.build_workbook(scn7(), student_name="=HYPERLINK(1)", student_id="S1")
    wb = load_workbook(io.BytesIO(data))
    assert wb[L.SHEET_START][L.START_NAME_CELL].data_type == "s"
