"""Help with designing the columns: the planner, what's new per plan, the heading check, the
diagram."""

import io

import pytest
from openpyxl import load_workbook

from aggplan import checker, columns, layout as L, workbook
from tests.helpers import SID, scn7

KINDS = ["chase", "level", "hybrid"]


@pytest.mark.parametrize("kind", KINDS)
def test_planner_covers_every_required_column_once(kind):
    listed = [c.concept for _, _, cols in columns.stages(kind) for c in cols]
    assert sorted(listed) == sorted(set(L.REQUIRED[kind]) | set(L.OPTIONAL[kind]))
    for _, _, cols in columns.stages(kind):
        for c in cols:
            assert c.question and c.why and c.comes_from, c.concept


@pytest.mark.parametrize("kind", KINDS)
def test_planner_ingredients_match_what_the_checker_requires(kind):
    """'Comes from' must name what the checker's left-out-input check requires, so the
    planner never tells a student less than the checker expects."""
    for concept, needs in checker._NEEDS.items():
        if concept not in columns.concepts(kind) or columns.is_decision(kind, concept):
            continue
        if concept in checker._CHASE_ONLY_NEEDS and kind != "chase":
            continue
        ins = set(columns.inputs(kind, concept))
        for alts in needs:
            want = {("Total Production" if kind == "hybrid" else "Regular Production")
                    if a == "Production" else a for a in alts}
            assert ins & want, (kind, concept, alts, ins)


def test_used_by_is_the_reverse_of_comes_from_and_uses_the_plans_titles():
    assert "Production" in columns.used_by("chase", "Beginning Inventory")
    assert "Hiring Cost" in columns.used_by("chase", "Hires")
    assert "next month's Beginning Inventory" in columns.used_by("chase", "Ending Inventory")
    assert "Total Production" in columns.used_by("hybrid", "Overtime Production")


def test_decisions_and_level_planning_cells_are_explained():
    assert "decision" in columns.comes_from("hybrid", "Workers").lower()
    assert "planning cell" in columns.comes_from("level", "Regular Production")
    assert "rounded UP" in columns.comes_from("chase", "Workers", whole_workers=True)
    assert "rounded" not in columns.comes_from("chase", "Workers", whole_workers=False)


def test_whats_new_does_not_give_the_new_columns_away():
    assert "14 questions" in columns.whats_new("chase")
    assert "Same 14" in columns.whats_new("level").replace("same", "Same")
    new = columns.whats_new("hybrid")
    assert "5 new questions" in new
    for c in ("Overtime Production", "Subcontract Production", "Total Production"):
        assert c not in new


def test_one_heading_at_a_time():
    found = {}
    r = columns.check_one("chase", "Starting Stock", found)
    assert r["status"] == "ok" and r["name"] == "Beginning Inventory"
    assert r["why"] and r["question"]                     # why it matters, once found
    found[r["concept"]] = "Starting Stock"
    assert columns.check_one("chase", "Opening inventory", found)["status"] == "repeat"
    assert columns.check_one("chase", "stuff", found) == {"status": "unknown"}   # no names
    assert columns.check_one("chase", "helper: notes", found)["status"] == "helper"
    assert columns.check_one("chase", "Overtime", found) == {"status": "other_plan",
                                                             "plans": ["hybrid"]}
    assert columns.check_one("hybrid", "Overtime", found)["status"] == "ok"
    assert columns.check_one("chase", "", found)["status"] == "empty"
    # a bare "Inventory" is read as the END of the month, and says so
    assert "end" in columns.check_one("chase", "Inventory", found)["note"]
    assert not columns.check_one("chase", "Closing stock", found)["note"]


def test_support_grows_with_each_miss():
    s1, s2, s3 = (columns.support(n) for n in (1, 2, 3))
    assert len({s1, s2, s3}) == 3 and len(s1) < len(s3)
    assert columns.support(9) == s3
    assert "inventory" not in s1.lower() and "inventory" in s3.lower()


def test_open_questions_list_questions_not_column_names():
    oq = columns.open_questions("hybrid", {"Beginning Inventory": "x"})
    flat = [(q, tag) for _, qs in oq for q, tag in qs]
    assert all("Inventory" not in q for q, _ in flat)
    assert ("How many bottles did you make on overtime?", "new") in flat
    assert ("How many bottles could your workforce make on regular time?", "optional") in flat
    assert "How many bottles are on the shelf when the month starts?" not in [q for q, _ in flat]


@pytest.mark.parametrize("kind", KINDS)
def test_diagram_is_valid_dot_with_every_column(kind):
    dot = columns.diagram(kind)
    assert dot.startswith("digraph") and dot.rstrip().endswith("}")
    assert dot.count("{") == dot.count("}")
    for c in columns.concepts(kind):
        assert f'"{c}"' in dot
    assert '"Ending Inventory" -> "Holding Cost"' in dot
    if kind == "level":
        assert '"Planning cells" -> "Regular Production"' in dot


def test_diagram_hides_columns_not_found_yet():
    dot = columns.diagram("chase", {"Workers": "Crew"})
    assert '"Workers" [fillcolor' in dot and 'label="Workers' in dot
    assert '"Hires" [fillcolor="#FFFFFF", label="?"' in dot


def test_start_here_has_no_column_planner():
    ws = load_workbook(io.BytesIO(workbook.build_workbook(scn7(), student_name="x",
                                                          student_id=SID)))[L.SHEET_START]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "My heading" not in values and not ws.row_breaks.brk


def test_order_advice_flags_columns_placed_before_what_they_use():
    good = columns.order_advice("chase", columns.concepts("chase"))
    assert good["ok"] and not good["problems"]
    assert columns.order_advice("chase", L.REQUIRED["chase"])["ok"]       # the usual order
    bad = columns.order_advice("chase", ["Total Cost"] + [c for c in L.REQUIRED["chase"]
                                                          if c != "Total Cost"])
    assert not bad["ok"]
    assert ("Total Cost", "Labor Cost") in bad["problems"]
    assert bad["recommended"][0] == "Beginning Inventory" and bad["recommended"][-1] == "Total Cost"
    # hires before workers: hires uses workers
    o = ["Hires", "Workers"]
    assert columns.order_advice("chase", o)["problems"] == [("Hires", "Workers")]
    # hybrid decisions use nothing, so typing them in any order is fine
    assert columns.order_advice("hybrid", ["Overtime Production", "Workers"])["ok"]
