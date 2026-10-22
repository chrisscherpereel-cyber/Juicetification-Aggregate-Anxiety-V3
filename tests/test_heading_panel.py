"""The 'Design your columns' panel in the app: one heading at a time, help that grows with
each miss, and a column's name and reason shown only once it is found."""

import os

import pytest

st_testing = pytest.importorskip("streamlit.testing.v1")
APP = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "app.py")


def chase_page():
    at = st_testing.AppTest.from_file(APP, default_timeout=60)
    at.query_params["seed"] = "7"
    at.run()
    at.text_input[0].input("Ann")
    at.text_input[1].input("S1")
    next(b for b in at.button if b.label == "Start").click()
    at.run()
    at.session_state["step"] = "chase"
    return at.run()


def check(at, heading):
    next(t for t in at.text_input if t.label == "Check one heading").input(heading)
    next(b for b in at.button if b.label == "Check").click()
    return at.run()


def texts(at):
    return " ".join([m.value for m in at.markdown] + [m.value for m in at.caption]
                    + [m.value for m in at.success] + [m.value for m in at.warning]
                    + [m.value for m in at.info])


def test_nothing_is_given_away_before_it_is_found():
    at = chase_page()
    assert not at.exception
    body = texts(at)
    assert "How many bottles are on the shelf when the month starts?" in body
    assert "Beginning Inventory" not in body and "Ending Backorders" not in body
    assert "0 of 14 columns found" in " ".join(p.value.get("text", "") if isinstance(p.value, dict)
                                                else str(p.proto.text) for p in at.get("progress"))


def test_misses_get_more_help_and_a_find_explains_why():
    at = chase_page()
    at = check(at, "stuff")
    w1 = at.warning[-1].value
    assert "isn't recognised" in w1 and "Beginning" not in w1
    at = check(at, "more stuff")
    assert "open questions" in at.warning[-1].value
    assert any("Try this one" in c.value for c in at.caption)
    at = check(at, "Starting Stock")
    ok = at.success[-1].value
    assert "Beginning Inventory" in ok and "Why it matters" in ok
    assert at.session_state["pieces"]["chase"]["found"] == {"Beginning Inventory": "Starting Stock"}
    assert at.session_state["pieces"]["chase"]["head_misses"] == 2
    assert "1 of 14 columns found" in str(at.get("progress")[0].proto)
    at = check(at, "Opening inventory")
    assert "same question" in at.info[-1].value
    assert "How many bottles are on the shelf when the month starts?" not in texts(at)


def test_newest_find_first_and_diagram_above_the_lists():
    at = chase_page()
    for h in ("Starting Stock", "Crew", "Hired"):
        at = check(at, h)
    found = [m.value for m in at.markdown if m.value.startswith("✔ **")]
    assert [f.split("**")[1] for f in found] == ["Hires", "Workers", "Beginning Inventory"]
    assert any("Still to find" in m.value for m in at.markdown)
    found_box = next(e for e in at.expander if e.label.startswith("Found so far (3)"))
    assert not found_box.proto.expanded                       # closed until opened
    labels = [e.label for e in at.expander]
    assert "How they connect" in labels and "About this step" in labels
    assert not next(e for e in at.expander if e.label == "About this step").proto.expanded


CHASE_HEADINGS = ["Starting Stock", "Backlog at start", "Production", "Crew", "Hired", "Fired",
                  "Closing Stock", "Ending Backorders", "Wages", "Hiring Cost", "Layoff Cost",
                  "Holding Cost", "Backorder Cost", "Total Cost"]


def test_check_scrolls_the_reply_into_view():
    at = check(chase_page(), "Starting Stock")
    html = " ".join(str(e.proto) for e in at.get("iframe"))
    assert "aa-headcheck-chase" in html and "scrollIntoView" in html


def test_all_found_in_a_good_order_says_so():
    at = chase_page()
    for h in CHASE_HEADINGS:
        at = check(at, h)
    assert any("Every column is found" in s.value and "works" in s.value for s in at.success)


def test_all_found_in_a_poor_order_gets_a_recommended_order_and_why():
    at = chase_page()
    for h in ["Total Cost"] + CHASE_HEADINGS[:-1]:
        at = check(at, h)
    body = " ".join(m.value for m in at.markdown)
    assert "Every column is found" in body and "**Total Cost** comes before" in body
    assert "A better order" in body and "**C** Beginning Inventory" in body
    assert "Why it's better" in body and "Mistakes flow one way" in body


# ---- the workbook comes after the columns -----------------------------------------------------
def download_labels(at):
    return [d.label for d in at.get("download_button")]


def test_workbook_is_offered_with_the_headings_once_every_chase_column_is_found():
    import io
    from openpyxl import load_workbook
    at = chase_page()
    assert "Download the workbook anyway (no headings)" in download_labels(at)
    assert not any("with the headings" in d for d in download_labels(at))
    for h in ["Total Cost"] + CHASE_HEADINGS[:-1]:            # found in a poor order
        at = check(at, h)
    assert any("with the headings" in d for d in download_labels(at))
    assert "Download the workbook anyway (no headings)" not in download_labels(at)
    data = at.session_state["wb_head_bytes"]
    wb = load_workbook(io.BytesIO(data))
    for sheet in ("Chase Plan", "Level Plan"):               # recommended order, standard names
        row1 = [wb[sheet].cell(1, j).value for j in range(3, 17)]
        assert row1[0] == "Beginning Inventory" and row1[-1] == "Total Cost", (sheet, row1)
        assert row1[3] == "Workers" and row1[8] == "Labor Cost"
    why = next(e for e in at.expander if e.label == "Why the workbook uses standard headings")
    body = " ".join(m.value for m in why.markdown)
    assert "standard terms" in body and "| Starting Stock | Beginning Inventory |" in body
    assert wb["Chase Plan"]["C2"].value is None                # no formulas given


def test_level_panel_starts_with_the_chase_columns_found():
    at = chase_page()
    for h in CHASE_HEADINGS:
        at = check(at, h)
    at.session_state["pieces"]["chase"]["submitted"] = {"at": "t", "sha": "c", "file": "w",
                                                        "metrics": {}, "feasibility": []}
    at.session_state["step"] = "level"
    at = at.run()
    assert not at.exception
    assert "14 of 14 columns found" in str(at.get("progress")[0].proto)
    assert not any("with the headings" in d for d in download_labels(at))   # chase only


def test_downloading_early_is_noted_in_the_report():
    fitz = pytest.importorskip("fitz")
    from aggplan import report
    from tests.helpers import scn7
    pdf, _ = report.build_report(
        scn7(), student_name="Ann", student_id="S1", identity_label="x", final=True,
        heading="Final report",
        pieces=[{"title": "Chase plan", "status": "Submitted", "attempts": 1, "help": 0,
                 "note": "downloaded the workbook before finishing the column design."}])
    text = "".join(p.get_text() for p in fitz.open(stream=pdf, filetype="pdf"))
    assert "before finishing the column design" in text


# ---- start over with a new scenario --------------------------------------------------------
def sidebar_button(at, label):
    return next(b for b in at.sidebar.button if b.label == label)


def test_start_over_warns_then_gives_a_new_scenario_and_clears_the_work():
    at = st_testing.AppTest.from_file(APP, default_timeout=60).run()   # no fixed seed
    at.text_input[0].input("Ann")
    at.text_input[1].input("S1")
    next(b for b in at.button if b.label == "Start").click()
    at = at.run()
    at.session_state["step"] = "chase"
    at = at.run()
    at = check(at, "Starting Stock")
    at.session_state["pieces"]["chase"]["submitted"] = {"at": "t", "sha": "c", "file": "w",
                                                        "metrics": {}, "feasibility": []}
    at = at.run()
    seed0 = next(c.value for c in at.sidebar.markdown if "Scenario #" in c.value)
    sidebar_button(at, "↺ Start over with a new scenario").click()
    at = at.run()
    warn = " ".join(w.value for w in at.sidebar.warning)
    assert "Start over?" in warn and "can't be undone" in warn and "new scenario" in warn
    sidebar_button(at, "Cancel").click()                     # cancel changes nothing
    at = at.run()
    assert at.session_state["pieces"]["chase"]["found"] and at.session_state["restarts"] == 0
    sidebar_button(at, "↺ Start over with a new scenario").click()
    at = at.run()
    sidebar_button(at, "Yes, start over").click()
    at = at.run()
    assert not at.exception
    ss = at.session_state
    assert ss["restarts"] == 1 and ss["step"] == "start"
    assert ss["pieces"]["chase"]["found"] == {} and not ss["pieces"]["chase"]["submitted"]
    assert ss["restart_log"][0]["submitted"] == ["chase"]
    seed1 = next(c.value for c in at.sidebar.markdown if "Scenario #" in c.value)
    assert seed1 != seed0 and ss["student"]["sid"] == "S1"


def test_fixed_scenario_game_keeps_the_numbers_and_says_so():
    at = st_testing.AppTest.from_file(APP, default_timeout=60)
    at.query_params["seed"] = "7"
    at = at.run()
    at.text_input[0].input("Ann")
    at.text_input[1].input("S1")
    next(b for b in at.button if b.label == "Start").click()
    at = at.run()
    sidebar_button(at, "↺ Start over with a new scenario").click()
    at = at.run()
    assert "numbers stay the same" in " ".join(w.value for w in at.sidebar.warning)
    sidebar_button(at, "Yes, start over").click()
    at = at.run()
    assert any("Scenario #7" in c.value for c in at.sidebar.markdown)
