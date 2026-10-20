"""The performance changes must not change any answer: the fast reader agrees with a full
read, caches return the same thing, saves coalesce, and a receipt is built once."""

import io
import os
import threading
import time

import pytest
from openpyxl import load_workbook

from aggplan import checker, engine, layout as L, saving, workbook
from tests.helpers import SID, fixture, scn7

APP = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "app.py")


def _slow_read(data):
    """The original full-mode reader, kept here as the reference."""
    wf = load_workbook(io.BytesIO(data), data_only=False)
    wv = load_workbook(io.BytesIO(data), data_only=True)
    out = {}
    for name in wf.sheetnames:
        for rf, rv in zip(wf[name].iter_rows(max_row=checker.READ_ROWS, max_col=checker.READ_COLS),
                          wv[name].iter_rows(max_row=checker.READ_ROWS, max_col=checker.READ_COLS)):
            for cf, cv in zip(rf, rv):
                if not hasattr(cf, "coordinate"):
                    continue
                f = cf.value if (isinstance(cf.value, str) and cf.data_type == "f") else None
                v = cv.value if (cv.value is not None or f) else cf.value
                if f is not None or v is not None:
                    out[(name, cf.coordinate)] = (v, f)
    return out


@pytest.mark.parametrize("name", ["key_seed7.xlsx", "mistakes_seed7.xlsx", "shuffled_seed7.xlsx"])
def test_fast_reader_matches_full_read(name):
    data = fixture(name)
    checker._BOOK_CACHE.clear()
    book, err = checker.read_book(data)
    assert err is None
    fast = {(s, ref): (c.value, c.formula) for s, cells in book.sheets.items()
            for ref, c in cells.items()}
    assert fast == _slow_read(data)


def test_reader_reads_only_the_sheets_a_piece_needs():
    checker._BOOK_CACHE.clear()
    book, _ = checker.read_book(fixture("key_seed7.xlsx"), checker.PIECE_SHEETS["chase"])
    assert set(book.sheets) == {L.SHEET_CHASE}
    assert book.has_sheet(L.SHEET_SUMMARY) and book.meta["student_id"] == SID


def test_same_upload_is_parsed_once():
    data = fixture("key_seed7.xlsx")
    checker._BOOK_CACHE.clear()
    a, _ = checker.read_book(data, checker.PIECE_SHEETS["level"])
    b, _ = checker.read_book(data, checker.PIECE_SHEETS["level"])
    assert a is b


def test_cache_is_bounded():
    checker._BOOK_CACHE.clear()
    for i in range(checker._BOOK_CACHE_SIZE + 5):
        checker._BOOK_CACHE[("x", i)] = None
    checker.read_book(fixture("key_seed7.xlsx"), ())
    assert len(checker._BOOK_CACHE) <= checker._BOOK_CACHE_SIZE


def test_reference_plans_are_cached_but_safe_to_modify():
    s = scn7()
    a = engine.reference_plan(s, "chase")
    a.loc[:, "Total Cost"] = 0                       # a caller scribbling on its copy...
    assert engine.reference_plan(s, "chase")["Total Cost"].sum() == 426000   # ...changes nothing


def test_workbook_cache_returns_the_same_file_and_respects_settings():
    s = scn7()
    a = workbook.build_workbook(s, student_name="A", student_id="S1")
    assert workbook.build_workbook(s, student_name="A", student_id="S1") is a
    assert workbook.build_workbook(s, student_name="B", student_id="S1") != a
    no_hybrid = workbook.build_workbook(scn7(include_hybrid=False), student_name="A",
                                        student_id="S1")
    assert L.SHEET_HYBRID not in load_workbook(io.BytesIO(no_hybrid)).sheetnames


def test_template_styles_match_regular_styling():
    wb = load_workbook(io.BytesIO(workbook.build_workbook(scn7(), student_id="S9")))
    for ref in ("C1", "AD1", "C2", "AD13", "C14", "C17"):
        c = wb[L.SHEET_CHASE][ref]
        assert c.protection.locked is False and c.border.left.style == "thin", ref
    assert wb[L.SHEET_CHASE]["C1"].fill.fgColor.rgb.endswith("FCE4D6")
    assert wb[L.SHEET_CHASE]["C14"].fill.fgColor.rgb.endswith("BDD7EE")


def test_background_saves_coalesce_and_snapshot():
    calls, gate = [], threading.Event()

    def slow_save(state, **kw):
        gate.wait(2)
        calls.append((state, kw))

    state = {"n": 1}
    saving.save_later("stu", slow_save, state, progress=0.1)   # starts, waits on the gate
    time.sleep(0.05)
    for n in (2, 3, 4):                                          # queue while it runs
        state["n"] = n
        saving.save_later("stu", slow_save, state, progress=n / 10)
    state["n"] = 99                                              # later edits don't leak in
    gate.set()
    for _ in range(100):
        if saving.pending() == 0 and len(calls) >= 2:
            break
        time.sleep(0.02)
    assert [c[0]["n"] for c in calls] == [1, 4]                 # first, then only the newest


st_testing = pytest.importorskip("streamlit.testing.v1")


def test_receipt_is_built_once_per_submission(monkeypatch):
    from aggplan import report
    count = {"n": 0}
    real = report.build_report

    def counting(*a, **k):
        count["n"] += 1
        return real(*a, **k)

    monkeypatch.setattr(report, "build_report", counting)
    at = st_testing.AppTest.from_file(APP, default_timeout=60).run()
    at.text_input[0].input("Ann")
    at.text_input[1].input("S1")
    next(b for b in at.button if b.label == "Start").click()
    at = at.run()
    at.session_state["pieces"]["chase"]["submitted"] = {"at": "now", "sha": "x", "file": "f",
                                                        "metrics": {}, "feasibility": []}
    at.session_state["step"] = "chase"
    for _ in range(4):
        at.run()
    assert not at.exception and count["n"] == 1
