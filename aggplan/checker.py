"""Check ONE piece of an uploaded workbook: the chase plan, the level plan, the hybrid plan
or the Summary comparison.

The checker never evaluates a formula. It reads the file twice with openpyxl: once for the
formulas the student typed and once for the values Excel calculated and saved. On a plan
sheet it first works out which column is which from the student's own headings
(recognize.py), then grades every month of every recognised column.

Cell statuses:
  ok        right value, calculated by a formula
  typed     right value, but typed in as a number (when formulas are required)
  carried   the cell follows the rule, but an earlier cell it uses is wrong
  wrong     the value is wrong
  blank     nothing there yet
  error     Excel shows an error (#DIV/0!, #NAME? ...)
  text      words where a number belongs
  uncalc    a formula with no saved value (file not saved by a spreadsheet program)
  decision  a hybrid decision cell (the student's choice; not graded for value)

The checker reports WHAT is wrong; how much help to show is the app's job (hints.py).
"""

from __future__ import annotations

import hashlib
import io
import math
import re
import threading
import zipfile
from collections import OrderedDict
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

import pandas as pd
from openpyxl.utils import column_index_from_string, get_column_letter

from . import MODEL_VERSION, engine, layout as L
from .engine import Decision
from .recognize import ColumnMap, map_columns
from .scenario import Scenario
from .workbook import APP_ID

MAX_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED = 60 * 1024 * 1024
MAX_MEMBERS = 2000
READ_ROWS, READ_COLS = 40, L.LAST_STUDENT_COL + 2

PROBLEM = {"wrong", "typed", "error", "text", "uncalc", "carried"}
NOT_GRADED = {"decision", "helper"}
STATUS_LABEL = {
    "ok": "✔ Correct", "typed": "✎ Typed, not calculated", "carried": "↪ Follows an earlier mistake",
    "wrong": "✖ Not right yet", "blank": "○ Empty", "error": "⚠ Excel error",
    "text": "⚠ Text, not a number", "uncalc": "⚠ Not calculated", "decision": "◆ Your decision",
    "waiting": "… Waiting for your decisions", "helper": "· Not checked (helper column)",
}
STATUS_HELP = {
    "typed": "The number is right but it was typed in. Replace it with a formula so Excel "
             "calculates it.",
    "carried": "This cell follows the rule, but it uses an earlier cell that is wrong. Fix "
               "that cell first and this one will correct itself.",
    "blank": "Still to do.",
    "error": "Excel could not calculate this cell. See 'What an error means' on the Excel "
             "Help sheet.",
    "text": "This cell holds text. Enter a formula that starts with =.",
    "uncalc": "The formula has no saved result. Open the file in Excel, save it as .xlsx and "
              "upload again.",
}
PIECE_TITLES = {"chase": "Chase plan", "level": "Level plan", "hybrid": "Hybrid plan",
                "summary": "Compare & recommend"}
_ERR_RE = re.compile(r"^#(DIV/0!|NAME\?|REF!|VALUE!|N/A|NUM!|NULL!|SPILL!|CALC!)")
_REF_RE = re.compile(r"(?<![A-Z$])(\$?)([A-Z]{1,2})(\$?)(\d+)(?![\d(])")


# =========================================================================================
# Reading
# =========================================================================================
@dataclass
class Cell:
    value: object = None
    formula: Optional[str] = None

    @property
    def blank(self) -> bool:
        return self.formula is None and (self.value is None or
                                         (isinstance(self.value, str) and not self.value.strip()))


@dataclass
class Book:
    sheets: Dict[str, Dict[str, Cell]]
    meta: Dict[str, str]
    sha256: str
    names: Tuple[str, ...] = ()          # every sheet in the file (parsed or not)

    def has_sheet(self, name: str) -> bool:
        return name in self.sheets or name in self.names

    def cell(self, sheet: str, ref: str) -> Cell:
        return self.sheets.get(sheet, {}).get(ref, Cell())


def _container_problem(data: bytes) -> Optional[str]:
    if not data:
        return "The file is empty."
    if len(data) > MAX_BYTES:
        return "The file is larger than 5 MB. Upload the workbook you downloaded from the app."
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > MAX_MEMBERS or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
                return "The file is too large inside to be this workbook."
            if "xl/workbook.xml" not in {i.filename for i in infos}:
                return "This is not an Excel .xlsx workbook."
    except zipfile.BadZipFile:
        return ("This is not an Excel .xlsx workbook. If you used another program, save or "
                "download it as an Excel Workbook (.xlsx).")
    return None


PIECE_SHEETS = {"chase": (L.SHEET_CHASE,), "level": (L.SHEET_LEVEL,),
                "hybrid": (L.SHEET_HYBRID,),
                "summary": (L.SHEET_SUMMARY, L.SHEET_CHASE, L.SHEET_LEVEL, L.SHEET_HYBRID)}

# Parsed uploads, shared by every session in the process. A student usually uploads the same
# file to several steps (and the app reruns while it stays in the uploader), so a parsed
# file is reused instead of unzipped and parsed again. Bounded: oldest entries drop out.
_BOOK_CACHE: "OrderedDict[tuple, Book]" = OrderedDict()
_BOOK_CACHE_SIZE = 96
_BOOK_LOCK = threading.Lock()


def read_book(data: bytes, sheets: Optional[Sequence[str]] = None
              ) -> Tuple[Optional[Book], Optional[str]]:
    """(Book, None) or (None, a message the student can act on). Never raises.

    `sheets` limits parsing to the sheets a check needs (plus the hidden id sheet); None
    reads every sheet. openpyxl's read-only mode streams the XML instead of building a full
    styled workbook — about 7x less CPU per upload, which is what lets one Streamlit
    process absorb a whole class uploading at once."""
    problem = _container_problem(data)
    if problem:
        return None, problem
    sha = hashlib.sha256(data).hexdigest()
    ckey = (sha, None if sheets is None else tuple(sheets))
    with _BOOK_LOCK:
        if ckey in _BOOK_CACHE:
            _BOOK_CACHE.move_to_end(ckey)
            return _BOOK_CACHE[ckey], None
    try:
        from openpyxl import load_workbook
        wf = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        wv = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:                                   # noqa: BLE001
        return None, f"The workbook could not be opened ({type(e).__name__}). Save it again as .xlsx."
    try:
        names = tuple(wf.sheetnames)
        wanted = [n for n in names if sheets is None or n in sheets]
        parsed: Dict[str, Dict[str, Cell]] = {}
        for name in wanted:
            cells: Dict[str, Cell] = {}
            for row_f, row_v in zip(
                    wf[name].iter_rows(min_row=1, max_row=READ_ROWS, max_col=READ_COLS),
                    wv[name].iter_rows(min_row=1, max_row=READ_ROWS, max_col=READ_COLS)):
                for cf, cv in zip(row_f, row_v):
                    if not hasattr(cf, "coordinate"):
                        continue                               # empty / merged placeholder
                    raw, formula = cf.value, None
                    if isinstance(raw, str) and raw.startswith("=") and cf.data_type == "f":
                        formula = raw
                    elif type(raw).__name__ in ("ArrayFormula", "DataTableFormula"):
                        formula = getattr(raw, "text", None) or "="
                    val = cv.value
                    if formula is None and val is None and raw is not None:
                        val = raw
                    if formula is not None or val is not None:
                        cells[cf.coordinate] = Cell(val, formula)
            parsed[name] = cells
        meta: Dict[str, str] = {}
        if L.SHEET_META in names:
            for k, v in wv[L.SHEET_META].iter_rows(min_row=1, max_row=30, max_col=2,
                                                   values_only=True):
                if k:
                    meta[str(k)] = "" if v is None else str(v)
    except Exception as e:                                   # noqa: BLE001
        return None, f"The workbook could not be read ({type(e).__name__}). Save it again as .xlsx."
    finally:
        wf.close()
        wv.close()
    book = Book(parsed, meta, sha, names)
    with _BOOK_LOCK:
        _BOOK_CACHE[ckey] = book
        while len(_BOOK_CACHE) > _BOOK_CACHE_SIZE:
            _BOOK_CACHE.popitem(last=False)
    return book, None


# =========================================================================================
# Results
# =========================================================================================
@dataclass
class CellCheck:
    sheet: str
    ref: str
    item: str          # hint key: concept, "plan:<key>", "total", "summary:<metric>" ...
    label: str         # e.g. "Ending Inventory, March"
    row_label: str     # month / "Total" / "Planning" / summary row
    col_label: str     # grid column (e.g. "F · Closing Stock")
    status: str
    shown: str = ""
    diag: str = ""     # a specific observation about THIS cell (shown from hint level 2)

    @property
    def needs_work(self) -> bool:
        return self.status not in ("ok", "decision", "helper")


@dataclass
class ColumnNote:
    kind: str                     # missing | duplicate | unknown | guessed
    concept: str = ""
    letters: Tuple[str, ...] = ()
    heading: str = ""


@dataclass
class PieceResult:
    piece: str
    title: str
    fatal: List[str] = field(default_factory=list)
    columns: List[Tuple[str, str, Optional[str]]] = field(default_factory=list)  # letter, heading, read as
    notes: List[ColumnNote] = field(default_factory=list)
    cells: List[CellCheck] = field(default_factory=list)
    messages: List[str] = field(default_factory=list)
    feasibility: List[dict] = field(default_factory=list)
    metrics: Dict[str, Dict[str, float]] = field(default_factory=dict)   # plan -> summarize()
    sha256: str = ""
    recommendation: str = ""
    justification: str = ""
    words: int = 0
    min_words: int = 0
    must_be_feasible: bool = False
    plan: Optional[pd.DataFrame] = None      # the plan as the answer key computes it

    @property
    def readable(self) -> bool:
        return not self.fatal

    @property
    def graded(self) -> List[CellCheck]:
        return [c for c in self.cells if c.status not in NOT_GRADED]

    @property
    def n_ok(self) -> int:
        return sum(c.status == "ok" for c in self.graded)

    @property
    def n_total(self) -> int:
        return len(self.graded)

    @property
    def missing_columns(self) -> List[str]:
        return [n.concept for n in self.notes if n.kind == "missing"]

    @property
    def feasible(self) -> bool:
        return all(c["ok"] for c in self.feasibility) if self.feasibility else True

    @property
    def uncalculated(self) -> bool:
        return any(c.status == "uncalc" for c in self.cells)

    @property
    def complete(self) -> bool:
        return (self.readable and not any(n.kind == "missing" for n in self.notes)
                and bool(self.graded) and all(c.status == "ok" for c in self.graded)
                and not self.messages
                and (self.feasible or not self.must_be_feasible))

    def problem_items(self) -> Dict[str, List[CellCheck]]:
        """Items (hint keys) with a real mistake on this attempt. Carried and blank cells are
        excluded, so unfinished work never unlocks more help."""
        out: Dict[str, List[CellCheck]] = {}
        for c in self.cells:
            if c.status in PROBLEM - {"carried"}:
                out.setdefault(c.item, []).append(c)
        return out

    def todo_items(self) -> Dict[str, List[CellCheck]]:
        out: Dict[str, List[CellCheck]] = {}
        for c in self.cells:
            if c.status == "blank":
                out.setdefault(c.item, []).append(c)
        return out

    def counts(self) -> Dict[str, int]:
        out: Dict[str, int] = {}
        for c in self.graded:
            out[c.status] = out.get(c.status, 0) + 1
        return out


# =========================================================================================
# Helpers
# =========================================================================================
def _num(cell: Cell) -> Optional[float]:
    v = cell.value
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return None if (math.isnan(f) or math.isinf(f)) else f
    return None


def _shown(cell: Cell) -> str:
    v = cell.value
    if v is None:
        return "" if cell.formula is None else "(no value)"
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".") if not v.is_integer() else f"{v:,.0f}"
    if isinstance(v, int):
        return f"{v:,}"
    return str(v)[:40]


def _tol(item: str) -> float:
    return 0.01 if item in ("Workers", "level_workers", "Peak Workforce") else 0.5


def _close(a: Optional[float], b: Optional[float], tol: float) -> bool:
    return a is not None and b is not None and abs(a - b) <= max(tol, 1e-9 * abs(b))


def _f(*xs) -> bool:
    return all(x is not None for x in xs)


def _judge(cell: Cell, key: Optional[float], local: Optional[float], tol: float,
           require_formula: bool) -> str:
    if cell.blank:
        return "blank"
    if cell.formula is not None and cell.value is None:
        return "uncalc"
    if isinstance(cell.value, str):
        return "error" if _ERR_RE.match(cell.value.strip()) else "text"
    v = _num(cell)
    if v is None:
        return "text"
    if _close(v, key, tol):
        return "typed" if require_formula and cell.formula is None else "ok"
    if _close(v, local, tol):
        return "carried"
    return "wrong"


def _diag(kind: str, concept: str, i: int, cell: Cell, v, pv, scn: Scenario) -> str:
    """One specific observation about a wrong cell (never the answer)."""
    if isinstance(cell.value, str) and _ERR_RE.match(cell.value.strip()):
        return f"Excel shows {cell.value.strip()} here."
    f = (cell.formula or "").upper()
    if i > 0 and f:
        for m in _REF_RE.finditer(f):
            if int(m.group(4)) > L.TOTAL_ROW and not m.group(3):
                ref = f"{m.group(2)}{m.group(4)}"
                orig = f"${m.group(2)}${int(m.group(4)) - i}"
                return (f"This formula points at {ref}: filling down moved a reference that "
                        f"should stay put. Lock it with $ signs in January's formula "
                        f"(probably {orig}) and fill down again.")
    x = _num(cell)
    if x is None:
        return ""
    if concept == "Workers" and kind == "chase" and scn.whole_workers:
        pr, cap = v.get("Regular Production", [None] * 12)[i], pv.get("capacity")
        if _f(pr, cap) and cap:
            raw = pr / cap
            if abs(x - raw) < 0.01 and abs(raw - round(raw)) > 0.01:
                return "This is a fraction of a worker — you can't employ part of a person."
            if abs(x - math.floor(raw)) < 0.01 and raw > math.floor(raw) + 1e-9:
                return "Rounding down leaves you short of capacity."
    if concept in ("Ending Inventory", "Ending Backorders", "Hires", "Layoffs",
                   "Regular Production") and x < -1e-9:
        return "This can't be negative."
    if concept in ("Hires", "Layoffs", "Beginning Inventory", "Beginning Backorders") and i == 0:
        return "January is a special case: there is no previous month in the table."
    return ""


# =========================================================================================
# Formula references: circular references, and formulas that point at the wrong column
# =========================================================================================
_OTHER_SHEET_RE = re.compile(r"(?:'[^']+'|[A-Za-z_][\w.]*)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?")
_RANGE_RE = re.compile(r"(?<![A-Z$])\$?([A-Z]{1,2})\$?(\d+):\$?([A-Z]{1,2})\$?(\d+)")
_STRING_RE = re.compile(r'"[^"]*"')


def formula_refs(formula: str) -> List[Tuple[int, int]]:
    """(column index, row) of every cell on the SAME sheet a formula uses; ranges expanded
    within the area the checker reads. References to other sheets and text in quotes are
    ignored."""
    f = _STRING_RE.sub("", (formula or "").upper())
    f = _OTHER_SHEET_RE.sub("", f)
    out: List[Tuple[int, int]] = []

    def add_range(m):
        c1, c2 = sorted((column_index_from_string(m.group(1)), column_index_from_string(m.group(3))))
        r1, r2 = sorted((int(m.group(2)), int(m.group(4))))
        for c in range(c1, min(c2, READ_COLS) + 1):
            for r in range(r1, min(r2, READ_ROWS) + 1):
                out.append((c, r))
        return " "

    f = _RANGE_RE.sub(add_range, f)
    for m in _REF_RE.finditer(f):
        out.append((column_index_from_string(m.group(2)), int(m.group(4))))
    return out


def circular_groups(book: Book, sheet: str) -> List[List[str]]:
    """Groups of cells on one sheet whose formulas depend on each other in a loop (Excel's
    "circular reference"). Excel can't calculate them, so it saves 0 or stale values."""
    graph: Dict[str, List[str]] = {}
    for ref, cell in book.sheets.get(sheet, {}).items():
        if cell.formula:
            graph[ref] = [f"{get_column_letter(c)}{r}" for c, r in formula_refs(cell.formula)]
    index: Dict[str, int] = {}
    low: Dict[str, int] = {}
    stack: List[str] = []
    on: set = set()
    groups: List[List[str]] = []
    counter = [0]
    for start in graph:                       # Tarjan's strongly connected components
        if start in index:
            continue
        work = [(start, iter(graph.get(start, ())))]
        index[start] = low[start] = counter[0]
        counter[0] += 1
        stack.append(start)
        on.add(start)
        while work:
            node, it = work[-1]
            nxt = next(it, None)
            if nxt is not None:
                if nxt not in graph:
                    continue                  # a value, not a formula: can't loop
                if nxt not in index:
                    index[nxt] = low[nxt] = counter[0]
                    counter[0] += 1
                    stack.append(nxt)
                    on.add(nxt)
                    work.append((nxt, iter(graph[nxt])))
                elif nxt in on:
                    low[node] = min(low[node], index[nxt])
                continue
            work.pop()
            if work:
                low[work[-1][0]] = min(low[work[-1][0]], low[node])
            if low[node] == index[node]:
                comp = []
                while True:
                    x = stack.pop()
                    on.discard(x)
                    comp.append(x)
                    if x == node:
                        break
                if len(comp) > 1 or node in graph.get(node, ()):
                    groups.append(sorted(comp, key=lambda r: (int(re.sub(r"\D", "", r)),
                                                              column_index_from_string(
                                                                  re.sub(r"\d", "", r)))))
    return groups


def _circular_message(groups: List[List[str]], hybrid: bool) -> str:
    cells = ", ".join(groups[0][:6]) + (" …" if len(groups[0]) > 6 else "")
    more = f" (and {len(groups) - 1} more loop{'s' if len(groups) > 2 else ''})" if len(groups) > 1 else ""
    msg = (f"**Circular reference:** {cells}{more} use each other's results in a loop, so "
           "Excel can't calculate them (it shows 0 or an old number). Follow the formulas "
           "from one of these cells until you get back to it, and break the loop.")
    if hybrid:
        msg += (" On the hybrid sheet this usually means a green decision cell (Workers, "
                "Regular Production, Overtime, Subcontract) has a formula that uses another "
                "column calculated from it. Type your decisions as numbers, or base a formula "
                "only on fixed numbers such as Demand or a helper column with your starting "
                "plan.")
    return msg


# What each column may use from the SAME month (a formula may also use anything those use).
_USES: Dict[str, Tuple[str, ...]] = {
    "Regular Production": ("Beginning Inventory", "Beginning Backorders"),
    "Workers": ("Regular Production",),
    "Regular Capacity": ("Workers",),
    "Total Production": ("Regular Production", "Overtime Production", "Subcontract Production"),
    "Hires": ("Workers", "Regular Capacity"),
    "Layoffs": ("Workers", "Regular Capacity"),
    "Ending Inventory": ("Beginning Inventory", "Beginning Backorders", "Regular Production",
                         "Total Production", "Ending Backorders"),
    "Ending Backorders": ("Beginning Inventory", "Beginning Backorders", "Regular Production",
                          "Total Production", "Ending Inventory"),
    "Labor Cost": ("Workers", "Regular Capacity"),
    "Hiring Cost": ("Hires",), "Layoff Cost": ("Layoffs",),
    "Overtime Cost": ("Overtime Production",), "Subcontract Cost": ("Subcontract Production",),
    "Holding Cost": ("Ending Inventory",), "Backorder Cost": ("Ending Backorders",),
    "Total Cost": tuple(c for c in L.COST_CONCEPTS if c != "Total Cost"),
}
# ... and from the PREVIOUS month.
_USES_PREV: Dict[str, Tuple[str, ...]] = {
    "Beginning Inventory": ("Ending Inventory",), "Beginning Backorders": ("Ending Backorders",),
    "Hires": ("Workers", "Regular Capacity"), "Layoffs": ("Workers", "Regular Capacity"),
}


def _closure(concept: str, kind: str) -> set:
    seen, todo = set(), [concept]
    while todo:
        c = todo.pop()
        if c in seen:
            continue
        seen.add(c)
        if kind == "hybrid" and c in L.DECISIONS:
            continue                          # the student's own numbers: nothing upstream
        todo.extend(_USES.get(c, ()))
    return seen


def _allowed(concept: str, kind: str) -> Tuple[set, set]:
    same = _closure(concept, kind)
    prev: set = set()
    for c in same:
        for q in _USES_PREV.get(c, ()):
            prev |= _closure(q, kind)
    return same, prev


def wrong_column_refs(kind: str, concept: str, formula: str, row: int,
                      concept_of: Dict[str, str]) -> List[Tuple[str, str]]:
    """Cells this formula uses that this column has no business using — e.g. Hires reading
    the Overtime column. The value can still look right by coincidence (both 0), so this is
    checked even when the number matches. Returns [(cell, the concept that column holds)]."""
    same, prev = _allowed(concept, kind)
    out = []
    for c, r in formula_refs(formula):
        letter = get_column_letter(c)
        other = concept_of.get(letter)
        if other is None or other == concept:
            continue                          # Month, Demand, helper columns: fine
        if r == row and other not in same:
            out.append((f"{letter}{r}", other))
        elif r == row - 1 and r >= L.FIRST_ROW and other not in prev and other not in same:
            out.append((f"{letter}{r}", other))
    return out


# =========================================================================================
# Plan sheets
# =========================================================================================
def _headings(book: Book, sheet: str) -> Dict[int, str]:
    out = {}
    for j in range(L.FIRST_STUDENT_COL, L.LAST_STUDENT_COL + 1):
        v = book.cell(sheet, f"{get_column_letter(j)}1").value
        if v is not None and str(v).strip():
            out[j] = str(v).strip()
    return out


def _col_values(book: Book, sheet: str, letter: str) -> List[Optional[float]]:
    return [_num(book.cell(sheet, f"{letter}{L.month_row(i)}")) for i in range(L.N_MONTHS)]


def _planning_values(book: Book, kind: str) -> Dict[str, Optional[float]]:
    sheet = L.PLAN_SHEETS[kind]
    return {pc.key: _num(book.cell(sheet, L.planning_ref(kind, pc.key))) for pc in L.PLANNING[kind]}


def _hybrid_decision(book: Book, cm: ColumnMap) -> Optional[Decision]:
    sheet = L.SHEET_HYBRID
    if "Workers" not in cm.by_concept or "Regular Production" not in cm.by_concept:
        return None
    w = _col_values(book, sheet, cm.by_concept["Workers"])
    r = _col_values(book, sheet, cm.by_concept["Regular Production"])
    ot = (_col_values(book, sheet, cm.by_concept["Overtime Production"])
          if "Overtime Production" in cm.by_concept else [0.0] * 12)
    sub = (_col_values(book, sheet, cm.by_concept["Subcontract Production"])
           if "Subcontract Production" in cm.by_concept else [0.0] * 12)
    ot = [x or 0.0 for x in ot]
    sub = [x or 0.0 for x in sub]
    if not (_f(*w) and _f(*r)) or any(x < 0 for x in w + r + ot + sub):
        return None
    return Decision(w, r, ot, sub)


@dataclass
class PlanRead:
    cm: ColumnMap
    values: Dict[str, List[Optional[float]]]
    planning: Dict[str, Optional[float]]
    key: Optional[pd.DataFrame]


def read_plan(book: Book, scn: Scenario, kind: str) -> PlanRead:
    """Map the student's columns and read their numbers; build the answer key (for the
    hybrid: from the student's own decisions)."""
    sheet = L.PLAN_SHEETS[kind]
    heads = _headings(book, sheet)
    if kind == "hybrid":
        dec = _hybrid_decision(book, map_columns(heads, kind))
        key = engine.simulate(scn, dec) if dec else None
    else:
        key = engine.reference_plan(scn, kind)

    def guess(concept: str, letter: str) -> bool:
        if key is None or concept not in key:
            return False
        vals = _col_values(book, sheet, letter)
        hits = sum(_close(a, float(b), _tol(concept)) for a, b in zip(vals, key[concept]))
        return hits >= 10

    cm = map_columns(heads, kind, guess)
    values = {c: _col_values(book, sheet, letter) for c, letter in cm.by_concept.items()}
    return PlanRead(cm, values, _planning_values(book, kind), key)


def _local(kind: str, col: str, i: int, v: Dict[str, List[Optional[float]]],
           pv: Dict[str, Optional[float]], scn: Scenario) -> Optional[float]:
    """What this cell SHOULD show given the student's own upstream cells."""
    p = scn.params
    d = float(scn.demand[i])
    none = [None] * L.N_MONTHS
    g = lambda name: v.get(name, none)[i]                          # noqa: E731
    prev = lambda name: v.get(name, none)[i - 1]                   # noqa: E731
    prod_col = "Total Production" if kind == "hybrid" else "Regular Production"
    if col == "Beginning Inventory":
        return float(p["beginning_inventory"]) if i == 0 else prev("Ending Inventory")
    if col == "Beginning Backorders":
        return float(p.get("beginning_backorders", 0)) if i == 0 else prev("Ending Backorders")
    if col == "Regular Production":
        if kind == "chase":
            t, bi, bb = pv.get("target"), g("Beginning Inventory"), g("Beginning Backorders")
            return max(0.0, d + t - bi + bb) if _f(t, bi, bb) else None
        return pv.get("level_rate") if kind == "level" else None
    if col == "Workers":
        if kind == "chase":
            cap, pr = pv.get("capacity"), g("Regular Production")
            if not (_f(cap, pr) and cap):
                return None
            return float(math.ceil(pr / cap - 1e-9)) if scn.whole_workers else pr / cap
        return pv.get("level_workers") if kind == "level" else None
    if col == "Regular Capacity":
        cap, w = pv.get("capacity"), g("Workers")
        return w * cap if _f(cap, w) else None
    if col == "Total Production":
        parts = [g("Regular Production"), g("Overtime Production") or 0.0,
                 g("Subcontract Production") or 0.0]
        return sum(parts) if _f(*parts) else None
    if col in ("Hires", "Layoffs"):
        w = g("Workers")
        wp = float(p["starting_workforce"]) if i == 0 else prev("Workers")
        if not _f(w, wp):
            return None
        return max(0.0, w - wp) if col == "Hires" else max(0.0, wp - w)
    if col in ("Ending Inventory", "Ending Backorders"):
        bi, bb, pr = g("Beginning Inventory"), g("Beginning Backorders"), g(prod_col)
        if not _f(bi, bb, pr):
            return None
        return max(0.0, bi + pr - bb - d) if col == "Ending Inventory" else max(0.0, d + bb - bi - pr)
    rate_of = {"Labor Cost": ("Workers", "regular_labor_cost"), "Hiring Cost": ("Hires", "hiring_cost"),
               "Layoff Cost": ("Layoffs", "layoff_cost"),
               "Overtime Cost": ("Overtime Production", "overtime_cost"),
               "Subcontract Cost": ("Subcontract Production", "subcontract_cost"),
               "Holding Cost": ("Ending Inventory", "holding_cost"),
               "Backorder Cost": ("Ending Backorders", "backorder_cost")}
    if col in rate_of:
        q, k = rate_of[col]
        x = g(q)
        if x is None and q in ("Overtime Production", "Subcontract Production"):
            x = 0.0
        return x * float(p[k]) if x is not None else None
    if col == "Total Cost":
        parts = [g(c) for c in L.COST_CONCEPTS[:-1] if c in v]
        return sum(parts) if parts and _f(*parts) else None
    return None


def _judge_decision(cell: Cell, concept: str) -> str:
    if cell.blank:
        return "decision" if concept in ("Overtime Production", "Subcontract Production") else "blank"
    if cell.formula is not None and cell.value is None:
        return "uncalc"
    x = _num(cell)
    if x is None:
        return "text"
    return "wrong" if x < 0 else "decision"


def _grade_plan(book: Book, scn: Scenario, kind: str, res: PieceResult) -> None:
    sheet = L.PLAN_SHEETS[kind]
    require = bool(scn.params.get("require_formulas", True))
    pr = read_plan(book, scn, kind)
    cm = pr.cm
    for letter, heading in cm.headings.items():
        res.columns.append((letter, heading, cm.read_as.get(letter)))
    for concept in cm.missing(kind):
        res.notes.append(ColumnNote("missing", concept))
    for concept, letters in cm.duplicates.items():
        res.notes.append(ColumnNote("duplicate", concept, tuple(letters)))
    for concept, letter in cm.guessed.items():
        res.notes.append(ColumnNote("guessed", concept, (letter,), cm.headings.get(letter, "")))
    for letter in cm.unknown:
        res.notes.append(ColumnNote("unknown", "", (letter,), cm.headings.get(letter, "")))

    # planning cells
    lr = engine.level_rate(scn)
    key_pv = {"capacity": scn.cap_month, "target": scn.target, "total_demand": scn.total_demand,
              "level_rate": lr, "level_workers": engine.workers_for(scn, lr),
              "ot_limit": scn.cap_month * float(scn["overtime_pct"])}
    pv = pr.planning
    for pc in L.PLANNING[kind]:
        ref = L.planning_ref(kind, pc.key)
        cell = book.cell(sheet, ref)
        local = None
        if pc.key == "level_rate" and _f(pv.get("total_demand"), pv.get("target")):
            local = (pv["total_demand"] + pv["target"] - float(scn["beginning_inventory"])
                     + float(scn.params.get("beginning_backorders", 0))) / 12
        elif pc.key == "level_workers" and _f(pv.get("level_rate"), pv.get("capacity")) and pv["capacity"]:
            need = pv["level_rate"] / pv["capacity"]
            local = float(math.ceil(need - 1e-9)) if scn.whole_workers else need
        elif pc.key == "ot_limit" and pv.get("capacity") is not None:
            local = pv["capacity"] * float(scn["overtime_pct"])
        st = _judge(cell, key_pv[pc.key], local, _tol(pc.key), require)
        res.cells.append(CellCheck(sheet, ref, f"plan:{pc.key}", pc.label, "Planning", pc.label,
                                   st, _shown(cell)))

    if pr.key is None:
        res.messages.append("Fill every month of your Workers and Regular Production decision "
                            "columns with numbers (0 or more) — then the rest of the sheet can "
                            "be checked.")
    order = sorted(cm.by_concept.items(), key=lambda kv: column_index_from_string(kv[1]))
    concept_of = {letter: concept for concept, letter in cm.by_concept.items()}
    all_letters = sorted(cm.headings, key=column_index_from_string)
    for i in range(L.N_MONTHS):
        for letter in all_letters:
            concept = concept_of.get(letter)
            ref = f"{letter}{L.month_row(i)}"
            cell = book.cell(sheet, ref)
            col_label = f"{letter} · {cm.headings[letter]}"
            label = f"{cm.headings[letter]}, {scn.months[i]}"
            diag = ""
            if concept is None:                      # helper / unrecognised: shown, not graded
                res.cells.append(CellCheck(sheet, ref, "helper", label, scn.months[i], col_label,
                                           "helper", _shown(cell)))
                continue
            if kind == "hybrid" and concept in L.DECISIONS:
                st = _judge_decision(cell, concept)
            elif pr.key is None:
                st = "waiting"
            else:
                st = _judge(cell, float(pr.key[concept].iloc[i]),
                            _local(kind, concept, i, pr.values, pv, scn), _tol(concept), require)
                diag = (_diag(kind, concept, i, cell, pr.values, pv, scn)
                        if st in ("wrong", "error") else "")
                bad = (wrong_column_refs(kind, concept, cell.formula, L.month_row(i), concept_of)
                       if st in ("ok", "carried") and cell.formula else [])
                if bad:
                    ref0, other = bad[0]
                    st = "wrong"
                    diag = (f"The number happens to fit, but this formula uses {ref0} "
                            f"({other}), which {concept} doesn't depend on. Check which "
                            "column it should point at.")
            res.cells.append(CellCheck(sheet, ref, concept, label, scn.months[i], col_label, st,
                                       _shown(cell), diag))
    loops = circular_groups(book, sheet)
    if loops:
        in_loop = {r for g in loops for r in g}
        for c in res.cells:
            if c.ref in in_loop and c.status != "helper":
                c.status, c.diag = "error", "Part of a circular reference."
        res.messages.append(_circular_message(loops, kind == "hybrid"))
    if pr.key is None:
        return
    for concept, letter in order:
        if concept not in L.COST_CONCEPTS and concept not in L.TOTAL_OPTIONAL:
            continue
        ref = f"{letter}{L.TOTAL_ROW}"
        cell = book.cell(sheet, ref)
        if cell.blank and concept not in L.COST_CONCEPTS:
            continue                                       # optional total, not built
        vals = pr.values[concept]
        if kind == "hybrid" and concept in ("Overtime Production", "Subcontract Production"):
            vals = [x or 0.0 for x in vals]
        st = _judge(cell, float(pr.key[concept].sum()), sum(vals) if _f(*vals) else None,
                    0.5, require)
        res.cells.append(CellCheck(sheet, ref, "total", f"{cm.headings[letter]}, year total",
                                   "Total", f"{letter} · {cm.headings[letter]}", st, _shown(cell)))
    res.metrics[kind] = engine.summarize(pr.key)
    res.plan = pr.key
    if kind == "hybrid":
        res.feasibility = engine.check_feasibility(pr.key, scn)
        res.must_be_feasible = True


# =========================================================================================
# Summary sheet
# =========================================================================================
_SUMMARY_OWN_HINTS = ("Total Hires", "Total Layoffs", "Peak Workforce", "Peak Inventory",
                      "Months With Backorders")


def _student_metric(pr: PlanRead, book: Book, kind: str, metric: str) -> Optional[float]:
    v = pr.values
    if metric == "Peak Workforce":
        return max(v["Workers"]) if "Workers" in v and _f(*v["Workers"]) else None
    if metric == "Peak Inventory":
        x = v.get("Ending Inventory")
        return max(x) if x and _f(*x) else None
    if metric == "Months With Backorders":
        b = v.get("Ending Backorders")
        return float(sum(x > 1e-6 for x in b)) if b and _f(*b) else None
    col = {"Total Hires": "Hires", "Total Layoffs": "Layoffs"}.get(metric, metric)
    letter = pr.cm.by_concept.get(col)
    if not letter:
        return None
    total = _num(book.cell(L.PLAN_SHEETS[kind], f"{letter}{L.TOTAL_ROW}"))
    if total is None and _f(*v[col]):
        total = sum(v[col])
    return total


def _grade_summary(book: Book, scn: Scenario, res: PieceResult) -> None:
    sheet = L.SHEET_SUMMARY
    require = bool(scn.params.get("require_formulas", True))
    kinds = ["chase", "level"] + (["hybrid"] if scn.params.get("include_hybrid", True) else [])
    reads = {k: read_plan(book, scn, k) for k in kinds}
    keys = {}
    for k in kinds:
        if reads[k].key is None:
            res.messages.append("Your hybrid plan's decisions can't be read any more — fix the "
                                "Hybrid Plan sheet first.")
            continue
        keys[k] = engine.summarize(reads[k].key)
    res.metrics = keys
    for row in L.SUMMARY_ROWS:
        r = L.summary_row(row.metric)
        for kind in kinds:
            if (row.hybrid_only and kind != "hybrid") or kind not in keys:
                continue
            ref = f"{L.SUMMARY_PLAN_COL[kind]}{r}"
            cell = book.cell(sheet, ref)
            item = "summary:" + (row.metric if row.metric in _SUMMARY_OWN_HINTS else "link")
            st = _judge(cell, keys[kind][row.metric],
                        _student_metric(reads[kind], book, kind, row.metric), _tol(row.metric),
                        require)
            res.cells.append(CellCheck(sheet, ref, item, f"{row.label} ({L.PLAN_LABEL[kind]})",
                                       row.label, L.PLAN_LABEL[kind], st, _shown(cell)))
    tc = L.summary_row("Total Cost")
    kc, kl = keys["chase"]["Total Cost"], keys["level"]["Total Cost"]
    sc, sl = _num(book.cell(sheet, f"B{tc}")), _num(book.cell(sheet, f"C{tc}"))

    cell = book.cell(sheet, L.LOWER_COST_CELL)
    want = {"chase", "level"} if abs(kc - kl) < 0.5 else ({"chase"} if kc < kl else {"level"})
    local = None
    if _f(sc, sl):
        local = {"chase", "level"} if abs(sc - sl) < 0.5 else ({"chase"} if sc < sl else {"level"})
    uses_totals = True
    if cell.blank:
        st = "blank"
    elif cell.formula is not None and cell.value is None:
        st = "uncalc"
    else:
        got = str(cell.value).strip().lower()
        uses_totals = cell.formula is None or "!" in cell.formula or any(
            r == tc and c in (2, 3) for c, r in formula_refs(cell.formula))
        if _ERR_RE.match(got.upper()):
            st = "error"
        elif not uses_totals:
            st = "wrong"
        elif got in want:
            st = "typed" if require and cell.formula is None else "ok"
        elif local and got in local:
            st = "carried"
        else:
            st = "wrong"
    res.cells.append(CellCheck(sheet, L.LOWER_COST_CELL, "summary:lower", "Lower-cost plan",
                               "Lower-cost plan", "Answer", st, _shown(cell),
                               "" if uses_totals else
                               f"This formula doesn't compare the two TOTAL PLAN COST cells "
                               f"(B{tc} and C{tc}), so it won't follow your plans."))
    cell = book.cell(sheet, L.SAVINGS_CELL)
    st = _judge(cell, abs(kc - kl), abs(sc - sl) if _f(sc, sl) else None, 0.5, require)
    res.cells.append(CellCheck(sheet, L.SAVINGS_CELL, "summary:savings", "Savings",
                               "Savings", "Answer", st, _shown(cell)))

    allowed = [L.PLAN_LABEL[k] for k in kinds]
    cell = book.cell(sheet, L.RECOMMEND_CELL)
    rec = str(cell.value or "").strip()
    match = next((a for a in allowed if a.lower() == rec.lower()), "")
    res.recommendation = match
    st = "blank" if cell.blank else ("ok" if match else "wrong")
    res.cells.append(CellCheck(sheet, L.RECOMMEND_CELL, "summary:recommend", "Your recommendation",
                               "Recommendation", "Answer", st, rec))
    cell = book.cell(sheet, L.JUSTIFY_CELL)
    just = str(cell.value or "").strip() if not isinstance(cell.value, (int, float)) else ""
    words = len(re.findall(r"[A-Za-z0-9$%][\w$%.,'’-]*", just))
    need = int(scn.params.get("min_justification_words", 0) or 0)
    res.justification, res.words, res.min_words = just, words, need
    st = "blank" if not just else ("wrong" if words < need else "ok")
    res.cells.append(CellCheck(sheet, L.JUSTIFY_CELL, "summary:justify", "Justification",
                               "Justification", "Answer", st, f"{words} words"))


# =========================================================================================
# Public API
# =========================================================================================
def check_identity(book: Book, scn: Scenario, expected_sid: Optional[str]) -> List[str]:
    m = book.meta
    if m.get("app") != APP_ID:
        return ["This doesn't look like an Aggregate Anxiety workbook from this app (its hidden "
                "identification sheet is missing). Upload the workbook you downloaded, or "
                "download a fresh one and copy your work into it."]
    out = []
    if m.get("model_version", "").split(".")[0] != MODEL_VERSION.split(".")[0]:
        out.append("This workbook was made by a different version of the app. Download a fresh "
                   "workbook and copy your work into it.")
    if expected_sid and m.get("student_id", "") != expected_sid:
        out.append(f"This workbook belongs to student ID '{m.get('student_id', '')}', not "
                   f"'{expected_sid}'. Upload your own workbook.")
    pol = [(k, lab) for k, lab in (("safety_stock_policy", "safety stock"),
                                     ("worker_policy", "workforce"))
           if m.get(k) and m.get(k) != str(scn.params.get(k))]
    if pol and str(m.get("seed", "")) == str(scn.seed):
        made = ", ".join(f"{lab}: {m[k]}" for k, lab in pol)
        now = ", ".join(f"{lab}: {scn.params.get(k)}" for k, lab in pol)
        out.append(f"This workbook was made with different policies ({made}) from the ones now "
                   f"selected ({now}). Switch back on step 1, or download a new workbook there.")
        return out
    if str(m.get("seed", "")) != str(scn.seed) or m.get("config_hash") != scn.config_hash:
        out.append("This workbook was made for a different scenario than the one you are signed "
                   "in with (or your instructor changed the assignment). Download your workbook "
                   "again and copy your work into it.")
    return out


def check_book(book: Book, scn: Scenario, piece: str,
               expected_sid: Optional[str] = None) -> PieceResult:
    res = PieceResult(piece, PIECE_TITLES[piece], sha256=book.sha256)
    res.fatal = check_identity(book, scn, expected_sid)
    if res.fatal:
        return res
    sheet = L.SHEET_SUMMARY if piece == "summary" else L.PLAN_SHEETS[piece]
    if not book.has_sheet(sheet):
        res.fatal.append(f"The '{sheet}' sheet is missing or was renamed. Keep the sheet names "
                         "as they were.")
        return res
    for kind in (["chase", "level"] if piece == "summary" else [piece]):
        s = L.PLAN_SHEETS[kind]
        for i, d in enumerate(scn.demand):
            if not _close(_num(book.cell(s, f"B{L.month_row(i)}")), float(d), 1e-6):
                res.fatal.append(f"{s} B{L.month_row(i)}: the demand for {scn.months[i]} was "
                                 f"changed. Put back {d:,.0f} (given data must not change).")
                return res
    if piece == "summary":
        _grade_summary(book, scn, res)
    else:
        _grade_plan(book, scn, piece, res)
    return res


def check_piece(data: bytes, scn: Scenario, piece: str,
                expected_sid: Optional[str] = None) -> PieceResult:
    book, err = read_book(data, PIECE_SHEETS[piece])
    if err:
        return PieceResult(piece, PIECE_TITLES[piece], fatal=[err],
                           sha256=hashlib.sha256(data or b"").hexdigest())
    return check_book(book, scn, piece, expected_sid)


def grid(res: PieceResult) -> pd.DataFrame:
    """Cells for display: rows = months + Total, columns = the student's columns."""
    rows: Dict[str, Dict[str, CellCheck]] = {}
    order, cols = [], []
    for c in res.cells:
        if c.row_label == "Planning" or c.sheet == L.SHEET_SUMMARY:
            continue
        if c.row_label not in rows:
            rows[c.row_label] = {}
            order.append(c.row_label)
        if c.col_label not in cols:
            cols.append(c.col_label)
        rows[c.row_label][c.col_label] = c
    return pd.DataFrame([rows[r] for r in order], index=order, columns=cols)
